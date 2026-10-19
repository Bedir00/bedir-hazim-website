"""Export verified leads to XLSX + CSV.

usage: python3 export.py leads.json out_basename
leads.json: list of check results with 'verdict' = [category, reason, menu_link]
"""
import csv, json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REGION = {'Koeln': 'Raum Köln', 'Muenchen': 'Raum München', 'Hamburg': 'Raum Hamburg', 'Leipzig': 'Raum Leipzig'}
COLS = ['Nr', 'Restaurant', 'Kategorie', 'Was ist schlecht', 'Link zur Karte', 'Website', 'Straße', 'PLZ', 'Stadt',
        'Telefon', 'E-Mail', 'Küche', 'Geprüft am']

rows = json.load(open(sys.argv[1]))
out = sys.argv[2]
data = []
for i, r in enumerate(rows, 1):
    cat, why, link = r['verdict'][:3]
    data.append([i, r['name'], cat, why, link, r.get('final_url') or r['url'], r['street'], r['postcode'],
                 r['city'] or REGION.get(r['region'], r['region']), r['phone'], r['email'],
                 r['cuisine'].replace(';', ', ').replace('_', ' '), '19.10.2026'])

with open(out + '.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(COLS)
    w.writerows(data)

wb = Workbook()
ws = wb.active
ws.title = 'Leads'
ws.append(COLS)
for row in data:
    ws.append(row)
head = PatternFill('solid', fgColor='1F2937')
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = head
fills = {'Keine Online-Karte': 'FDE2E1', 'Nur PDF/Scan': 'FEF3C7', 'Schlechtes Design': 'E0E7FF'}
widths = [5, 30, 18, 60, 45, 35, 26, 8, 16, 18, 28, 16, 11]
for i, wdt in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = wdt
for row in ws.iter_rows(min_row=2):
    row[2].fill = PatternFill('solid', fgColor=fills.get(row[2].value, 'FFFFFF'))
    row[3].alignment = Alignment(wrap_text=True, vertical='top')
    for idx in (4, 5):
        if row[idx].value:
            row[idx].hyperlink = row[idx].value
            row[idx].font = Font(color='1D4ED8', underline='single')
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions
wb.save(out + '.xlsx')
print(len(data), 'rows ->', out + '.xlsx /.csv')
